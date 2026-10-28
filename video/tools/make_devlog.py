"""devlog.json → DEVLOG.xlsx (엑셀 일지) + DEVLOG.md (깃허브용) 만들기.

사용: python tools/make_devlog.py
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="FF8C28")
THIN = Side(style="thin", color="D0D0D0")
STATUS_FILLS = {"완료": "D9F2D9", "진행 중": "FFF2CC", "대기": "EEEEEE",
                "수정": "D9F2D9", "확인 대기": "FFF2CC"}


def style_sheet(ws, widths: list[int], header_row: int = 1) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=header_row):
        for c in row:
            c.font = Font(name=FONT, size=10, bold=c.row == header_row,
                          color="FFFFFF" if c.row == header_row else "000000")
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if c.row == header_row:
                c.fill = HEADER_FILL
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def add_table(ws, name: str, ref: str) -> None:
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)


def build_xlsx(data: dict, out: Path) -> None:
    wb = Workbook()

    # ---- 요약 ----
    ws = wb.active
    ws.title = "요약"
    n_road = len(data["roadmap"])
    n_ent = len(data["entries"])
    n_test = len(data["tests"])
    rows = [
        ("ReelCut 개발 일지", ""),
        ("", ""),
        ("로드맵 항목 수", f"=COUNTA('로드맵'!B2:B{n_road + 1})"),
        ("완료", f"=COUNTIF('로드맵'!E2:E{n_road + 1},\"완료\")"),
        ("진행 중", f"=COUNTIF('로드맵'!E2:E{n_road + 1},\"진행 중\")"),
        ("대기", f"=COUNTIF('로드맵'!E2:E{n_road + 1},\"대기\")"),
        ("진행률", "=IFERROR(B4/B3,0)"),
        ("", ""),
        ("기록 수", f"=COUNTA('개발 기록'!C2:C{n_ent + 1})"),
        ("사용자 제보 처리", f"=COUNTIF('개발 기록'!B2:B{n_ent + 1},\"사용자 제보\")"),
        ("최근 테스트 수", f"=INDEX('테스트'!B2:B{n_test + 1},{n_test})"),
        ("최근 테스트 통과", f"=INDEX('테스트'!C2:C{n_test + 1},{n_test})"),
        ("", ""),
        ("우선순위 기준", ""),
    ]
    rows += [(f"기준 {i}", c) for i, c in enumerate(data["criteria"], 1)]
    for r in rows:
        ws.append(r)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 70
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True, horizontal="left")
    ws["A1"].font = Font(name=FONT, size=16, bold=True)
    ws["A14"].font = Font(name=FONT, size=11, bold=True)
    ws["B7"].number_format = "0%"
    for r in range(3, 13):
        ws.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True)

    # ---- 로드맵 ----
    ws = wb.create_sheet("로드맵")
    ws.append(["순위", "항목", "기준", "이유", "상태", "완료일"])
    for r in data["roadmap"]:
        ws.append([r["rank"], r["item"], f"기준 {r['criterion']}", r["reason"], r["status"], r["done"]])
    style_sheet(ws, [6, 48, 9, 48, 10, 12])
    add_table(ws, "Roadmap", f"A1:F{n_road + 1}")
    rng = f"E2:E{n_road + 1}"
    for status, color in STATUS_FILLS.items():
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="equal", formula=[f'"{status}"'], fill=PatternFill("solid", fgColor=color)))

    # ---- 개발 기록 ----
    ws = wb.create_sheet("개발 기록")
    ws.append(["날짜", "구분", "항목", "내용", "결과", "커밋"])
    for e in data["entries"]:
        ws.append([e["date"], e["type"], e["title"], e["detail"], e["result"], e["commit"]])
    style_sheet(ws, [12, 11, 34, 70, 14, 10])
    add_table(ws, "Log", f"A1:F{n_ent + 1}")

    # ---- 테스트 ----
    ws = wb.create_sheet("테스트")
    ws.append(["날짜", "테스트 수", "통과", "통과율", "메모"])
    for i, t in enumerate(data["tests"], 2):
        ws.append([t["date"], t["total"], t["passed"], f"=IFERROR(C{i}/B{i},0)", t["note"]])
        ws.cell(row=i, column=4).number_format = "0%"
    style_sheet(ws, [12, 10, 8, 8, 50])
    add_table(ws, "Tests", f"A1:E{n_test + 1}")

    wb.save(out)


def build_md(data: dict, out: Path) -> None:
    lines = ["# ReelCut 개발 일지", "",
             "> 같은 내용을 엑셀로도 정리했어요: [`DEVLOG.xlsx`](DEVLOG.xlsx) (원본 데이터: `devlog.json`)", "",
             "## 우선순위 기준 (위가 먼저)", ""]
    lines += [f"{i}. {c}" for i, c in enumerate(data["criteria"], 1)]
    lines += ["", "## 로드맵", "", "| 순위 | 항목 | 기준 | 이유 | 상태 |", "|---|---|---|---|---|"]
    icon = {"완료": "✅", "진행 중": "🔧", "대기": "⏳"}
    for r in data["roadmap"]:
        lines.append(f"| {r['rank']} | {r['item']} | {r['criterion']} | {r['reason']} | "
                     f"{icon.get(r['status'], '')} {r['status']} |")
    lines += ["", "## 기록 (최신순)", ""]
    for e in reversed(data["entries"]):
        commit = f" (`{e['commit']}`)" if e["commit"] else ""
        lines.append(f"- **{e['date']} · {e['type']} · {e['title']}** — {e['detail']} → {e['result']}{commit}")
    t = data["tests"][-1]
    lines += ["", f"최근 테스트: {t['date']} · {t['passed']}/{t['total']} 통과 ({t['note']})", ""]
    out.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    data = json.loads((ROOT / "devlog.json").read_text(encoding="utf-8"))
    build_xlsx(data, ROOT / "DEVLOG.xlsx")
    build_md(data, ROOT / "DEVLOG.md")
    print("DEVLOG.xlsx / DEVLOG.md 갱신")


if __name__ == "__main__":
    main()
